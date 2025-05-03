# budget_app/app/utils.py
from io import BytesIO
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter, A4
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib import colors
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from datetime import datetime
from babel.numbers import format_currency

def generate_pdf_report(transactions, start_date=None, end_date=None):
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4,
                            leftMargin=0.75*inch,
                            rightMargin=0.75*inch,
                            topMargin=1*inch,
                            bottomMargin=1*inch)
    styles = getSampleStyleSheet()
    story = []

    # Titre
    title_str = "Rapport de Transactions"
    if start_date or end_date:
        period_str = "Période : "
        if start_date:
            period_str += f"du {start_date.strftime('%d/%m/%Y')} "
        if end_date:
            display_end_date = end_date.replace(hour=0, minute=0, second=0, microsecond=0)
            period_str += f"au {display_end_date.strftime('%d/%m/%Y')}"
        else:
             period_str += f"depuis le {start_date.strftime('%d/%m/%Y')}"
        title_str += f"\n{period_str}"

    story.append(Paragraph(title_str.replace('\n', '<br/>'), styles['h1']))
    story.append(Spacer(1, 0.2*inch))

    # Résumé
    total_income = sum(t.amount for t in transactions if t.type == 'income')
    total_expense = sum(t.amount for t in transactions if t.type == 'expense')
    balance = total_income - total_expense

    summary_text = f"<b>Total Revenus :</b> {format_currency(total_income, 'XOF', locale='fr_FR')}<br/>" \
                   f"<b>Total Dépenses :</b> {format_currency(total_expense, 'XOF', locale='fr_FR')}<br/>" \
                   f"<b>Solde :</b> {format_currency(balance, 'XOF', locale='fr_FR')}"
    story.append(Paragraph(summary_text, styles['Normal']))
    story.append(Spacer(1, 0.3*inch))

    # Tableau des transactions
    col_widths = [1.0*inch, 1.2*inch, 0.8*inch, 1.2*inch, 2.57*inch]

    data = [['Date', 'Montant (F CFA)', 'Type', 'Catégorie', 'Description']]
    for t in transactions:
        data.append([
            t.date.strftime('%d/%m/%Y'),
            format_currency(t.amount, 'XOF', locale='fr_FR'),
            'Revenu' if t.type == 'income' else 'Dépense',
            t.category.name if t.category else 'N/A',
            Paragraph(t.description or '', styles['Normal'])
        ])

    table = Table(data, colWidths=col_widths)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.darkblue),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, 0), 'CENTER'),
        ('ALIGN', (0, 1), (0, -1), 'CENTER'), # Date
        ('ALIGN', (1, 1), (1, -1), 'RIGHT'),  # Montant aligné à droite
        ('ALIGN', (2, 1), (2, -1), 'CENTER'), # Type
        ('ALIGN', (3, 1), (3, -1), 'LEFT'),   # Catégorie
        ('ALIGN', (4, 1), (4, -1), 'LEFT'),   # Description
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 11),
        ('FONTSIZE', (0, 1), (-1, -1), 10),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 10),
        ('GRID', (0, 0), (-1, -1), 1, colors.black),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.aliceblue])
    ]))

    story.append(table)
    story.append(Spacer(1, 0.5*inch))
    story.append(Paragraph(f"Généré le : {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}", styles['small']))

    doc.build(story)

    buffer.seek(0)
    return buffer

def generate_excel_report(transactions, start_date=None, end_date=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "Transactions"

    # Titre et période
    title_str = "Rapport de Transactions"
    ws.append([title_str])
    ws.merge_cells('A1:E1')
    ws['A1'].font = Font(bold=True, size=16)
    ws['A1'].alignment = Alignment(horizontal='center')

    if start_date or end_date:
        period_str = "Période : "
        if start_date:
            period_str += f"du {start_date.strftime('%d/%m/%Y')} "
        if end_date:
             display_end_date = end_date.replace(hour=0, minute=0, second=0, microsecond=0)
             period_str += f"au {display_end_date.strftime('%d/%m/%Y')}"
        else:
             period_str += f"depuis le {start_date.strftime('%d/%m/%Y')}"
        ws.append([period_str])
        ws.merge_cells('A2:E2')
        ws['A2'].font = Font(italic=True)
        ws['A2'].alignment = Alignment(horizontal='center')
        current_row = 3
    else:
        current_row = 2

    ws.append([]) # Ligne vide

    # En-têtes
    # MODIFICATION ICI (Optionnelle): Changement de l'en-tête Montant
    headers = ['Date', 'Montant (F CFA)', 'Type', 'Catégorie', 'Description']
    ws.append(headers)
    header_row_num = current_row + 1
    for col_idx, header in enumerate(headers, 1):
         cell = ws.cell(row=header_row_num, column=col_idx)
         cell.font = Font(bold=True)
         cell.alignment = Alignment(horizontal='center')


    # Données
    for t in transactions:
        ws.append([
            t.date.strftime('%Y-%m-%d'),
            t.amount,
            'Revenu' if t.type == 'income' else 'Dépense',
            t.category.name if t.category else 'N/A',
            t.description or ''
        ])
        last_row = ws.max_row
        # MODIFICATION ICI: Format numérique Excel pour XOF (sans décimales)
        ws[f'B{last_row}'].number_format = '#,##0 "F CFA"'
        ws[f'C{last_row}'].alignment = Alignment(horizontal='center')


    # Ajuster la largeur des colonnes
    ws.column_dimensions['A'].width = 12
    ws.column_dimensions['B'].width = 18 
    ws.column_dimensions['C'].width = 10
    ws.column_dimensions['D'].width = 20
    ws.column_dimensions['E'].width = 40

    # Sauvegarder dans un buffer mémoire
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer